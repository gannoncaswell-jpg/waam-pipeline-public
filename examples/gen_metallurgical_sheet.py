'Build the metallurgical bead-on-plate DOE capture workbook + section CSV.'


from __future__ import annotations

import argparse
import csv
import io
import sys
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from waam_pipeline import arcparams


WIRE_DIA_MM = arcparams.DEFAULT_WIRE_DIA_MM
IPM_TO_M_MIN = 0.0254


def expected_area_mm2(wfs_ipm: float, travel_mm_min: float) -> float:
    return round(arcparams.deposited_area_mm2(
        WIRE_DIA_MM, wfs_ipm * IPM_TO_M_MIN, travel_mm_min, electrode_efficiency=1.0), 3)


WFS_IPM = (180.0, 240.0, 300.0)
TRAVEL_MM_MIN = (250.0, 350.0, 450.0, 600.0)
CTWD_MM = 15.0
GAS = "90/10 Ar/CO2"
GAS_FLOW_L_MIN = 18.0
PROC = "PROC#4 (verified)"
DEPOSIT_LEN_MM = 120.0
INTERPASS_MAX_C = 120.0
SECTIONS = ("START", "MID", "STOP")


def beads() -> list[dict]:
    rows, n = [], 0
    for wfs in WFS_IPM:
        for ts in TRAVEL_MM_MIN:
            plate = n // 4 + 1
            n += 1
            rows.append(dict(bead_id=f"P{plate}-B{(n - 1) % 4 + 1}", plate=f"P{plate}",
                             wfs_ipm=wfs, travel_mm_min=ts))
    return rows


RUN_COLS = ["bead_id", "plate", "target_wfs_ipm", "target_travel_mm_min", "expected_area_mm2",
            "arc_mode", "target_ctwd_mm", "gas", "gas_flow_l_min", "deposit_len_mm", "interpass_max_c",
            "run_order", "actual_wfs_ipm", "actual_voltage_v", "actual_current_a",
            "start_temp_c", "interpass_c", "actual_travel_mm_min", "heat_input_j_mm",
            "operator", "run_date", "outcome", "run_notes"]


SECTION_COLS = ["section_id", "bead_id", "location", "cut_x_mm",
                "width_mm", "reinf_height_mm", "penetration_depth_mm",
                "reinf_area_mm2", "penetration_area_mm2", "toe_angle_l_deg", "toe_angle_r_deg",
                "expected_area_mm2", "area_vs_expected",
                "total_fused_area_mm2", "dilution_pct", "aspect_ratio_w_h",
                "parabola_area_mm2", "parabola_closure",
                "defects", "image_ref", "measured_by", "measure_date", "section_notes"]

DEFINITIONS = [
    ("Datum", "The ORIGINAL plate top surface, extended across the bead as a straight line. "
              "Everything above is reinforcement (added wire); everything below is penetration (remelted base)."),
    ("width_mm  (w)", "Bead width measured AT the datum line."),
    ("reinf_height_mm  (h)", "Max height of the deposit ABOVE the datum."),
    ("penetration_depth_mm  (p)", "Max depth of the fusion boundary BELOW the datum."),
    ("reinf_area_mm2  (A_r)", "Deposit cross-section area ABOVE the datum (planimeter in ImageJ). "
                              "This is the area that STACKS -- the model's primary input."),
    ("penetration_area_mm2  (A_p)", "Fused base-metal area BELOW the datum (planimeter). Remelt, not added mass."),
    ("toe_angle_*_deg", "Wetting/contact angle where the bead meets the plate, each side. "
                        "Low angle = good wetting = better overlap; high angle risks lack-of-fusion valleys."),
    ("dilution_pct", "= A_p / (A_r + A_p) x 100. Fraction of the bead that is remelted base metal."),
    ("parabola_area_mm2", "= (2/3) * w * h. The area of an ideal parabolic cap; compare to measured A_r."),
    ("parabola_closure", "= A_r / parabola_area. ~1.0 => the cap is well modelled as a parabola "
                         "(arcparams assumes A = 2/3 w h); far from 1 => use measured A_r directly."),
    ("SECTION LOCATIONS", "MID = steady state, cut at mid-length (feeds the geometry model). "
                          "START = in the arc-start region (hump). STOP = in the crater (valley). "
                          "START/STOP quantify the transient path planning must ramp out (Mehnen)."),
    ("expected_area_mm2", "PREDICTED by the model, not measured: arcparams.deposited_area_mm2(WFS, travel), "
                          "the pure mass balance wire_area*WFS/travel at eta=1.0. The ONE bead characteristic "
                          "the model predicts from setpoints -- and the only pre-filled expected value, because "
                          "the width/height split is empirical and arcparams refuses to guess it (this DOE "
                          "measures it). Blue cells = predicted by code; do not type in them."),
    ("area_vs_expected", "= measured A_r / expected_area_mm2. At MID this reads back the deposition efficiency "
                         "(~0.9-0.98). At START/STOP it shows the transient: >1 = start hump, <1 = crater."),
    ("heat_input_j_mm  (Q)", "= 0.8 * V * I / (travel mm/s), eta=0.8 for GMAW (Elsokaty Eq. 1). "
                             "Computed on Run Matrix from the as-run meter readings."),
    ("Model feed", "Deterministic, all through arcparams (docs/casting-slicing.md section 6): "
                   "expected area = deposited_area_mm2(); measured A_r / expected = retention efficiency; "
                   "A_r vs free_bead_area_mm2(w,h)=2/3 w h is the parabola check; stepover = "
                   "optimal_stepover_fraction(w,h) (Hu 2019 VOM); wall layer height = layer_height_for_area(A,T). "
                   "Measuring w,h here also retires the unmeasured 9.0/3.0 placeholders in clearance.py."),
]

PROTOCOL = [
    "1. CUT transverse (perpendicular to travel). MID at mid-length; if doing the transient, "
    "one section IN the start hump and one IN the crater. Bandsaw/abrasive/Wire-EDM, leave ~3-5 mm, "
    "keep coolant so you do not alter the microstructure. Label each piece with its section_id.",
    "2. MOUNT (hot compression or cold epoxy) with the transverse face down.",
    "3. GRIND progressively: 120 -> 240 -> 400 -> 600 -> 800 -> 1200 grit SiC, rotating 90 deg each step "
    "until the prior scratches are gone.",
    "4. POLISH: 3 um then 1 um (0.5 um alumina optional, per Elsokaty). Clean/dry between steps.",
    "5. ETCH: 2% Nital (2 mL HNO3 in 98 mL ethanol) -- ER70S-6 is low-carbon steel. SWAB 5-20 s in a fume "
    "hood, gloves + face shield, until the fusion boundary and HAZ show; rinse ethanol, warm-air dry. "
    "Nital is flammable/corrosive -- follow the SDS, never store mixed acid sealed.",
    "6. PHOTOGRAph the macro on a stereo/optical scope WITH A SCALE BAR (or stage micrometer) in frame; "
    "save the filename into image_ref.",
    "7. MEASURE in ImageJ: calibrate to the scale bar, trace the datum, then w, h, p, and planimeter "
    "A_r and A_p; read both toe angles. Enter into Section Capture; derived columns compute themselves.",
    "8. QA: reject a MID section with lack-of-fusion, a start/crater defect, or that is not steady state "
    "(compare to the two neighbours). Note any porosity/cracks/undercut in defects.",
]


def build_workbook(gen_utc: str):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    hdr_fill = PatternFill("solid", fgColor="1F3864")
    hdr_font = Font(bold=True, color="FFFFFF")
    derived_fill = PatternFill("solid", fgColor="E2EFDA")
    predicted_fill = PatternFill("solid", fgColor="DDEBF7")
    target_fill = PatternFill("solid", fgColor="FCE4D6")
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Border(*(Side(style="thin", color="BFBFBF"),) * 4)

    def header(ws, cols, note=None):
        r = 1
        if note:
            ws.cell(1, 1, note).font = Font(italic=True, color="808080")
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
            r = 2
        for c, name in enumerate(cols, 1):
            cell = ws.cell(r, c, name)
            cell.fill, cell.font, cell.border = hdr_fill, hdr_font, thin
            cell.alignment = wrap
        ws.freeze_panes = ws.cell(r + 1, 1)
        return r


    ws = wb.active
    ws.title = "Read Me"
    lines = [
        ("Metallurgical bead-on-plate DOE -- capture record sheet", True),
        (f"Generated {gen_utc}   |   cell: LC50S / AR1440   |   wire: ER70S-6 0.045 in", False),
        ("", False),
        ("PURPOSE: run a plate of beads at several parameter settings, section + macro-etch them,", False),
        ("and record the cross-section geometry that feeds the bead model (area, penetration,", False),
        ("dilution) plus the start/stop transient. Width-only, non-destructive lives in the", False),
        ("bead-width pack; this is its destructive complement. See docs/metallurgical-bead-doe.md.", False),
        ("", False),
        ("HOW TO USE:", True),
        ("  1. Run Matrix -- dial each bead's setpoints, then record the AS-RUN meter values.", False),
        ("     Orange cells are setpoints to dial; heat_input computes from the as-run readings.", False),
        ("  2. Section + macro-etch per the Macro Protocol tab.", False),
        ("  3. Section Capture -- one row per section; type the MEASURED cells (white).", False),
        ("     GREEN cells (dilution, parabola, etc.) are formulas -- do not type in them.", False),
        ("  4. Definitions defines every quantity + the datum. Diagrams: bead-section-locations.svg", False),
        ("     (where to cut) and bead-section-diagram.svg (what to measure).", False),
        ("", False),
        ("SEED DOE is a planning bracket, NOT an authorized schedule -- confirm WFS/trim against the", False),
        ("PROC#4 synergic curve on the pendant. The as-run columns are the record of truth.", False),
        ("", False),
        ("References: Elsokaty 2023 (ER70S-6 geometry), Mehnen 2011 (crossing/peak strategy),", False),
        ("Hu 2019 (multi-bead overlap), Ding 2016 (paths + bead model).", False),
    ]
    for i, (text, bold) in enumerate(lines, 1):
        cell = ws.cell(i, 1, text)
        cell.font = Font(bold=bold, size=13 if (bold and i == 1) else 11)
    ws.column_dimensions["A"].width = 100


    ws = wb.create_sheet("Run Matrix")
    hr = header(ws, RUN_COLS, "Orange = setpoint to dial.  heat_input_j_mm computes from as-run V, I, travel (eta=0.8).")
    col = {name: get_column_letter(i) for i, name in enumerate(RUN_COLS, 1)}
    for j, bead in enumerate(beads()):
        r = hr + 1 + j
        vals = {"bead_id": bead["bead_id"], "plate": bead["plate"],
                "target_wfs_ipm": bead["wfs_ipm"], "target_travel_mm_min": bead["travel_mm_min"],
                "expected_area_mm2": expected_area_mm2(bead["wfs_ipm"], bead["travel_mm_min"]),
                "arc_mode": PROC, "target_ctwd_mm": CTWD_MM, "gas": GAS,
                "gas_flow_l_min": GAS_FLOW_L_MIN, "deposit_len_mm": DEPOSIT_LEN_MM,
                "interpass_max_c": INTERPASS_MAX_C}
        for name in RUN_COLS:
            cell = ws.cell(r, RUN_COLS.index(name) + 1, vals.get(name))
            cell.border, cell.alignment = thin, wrap
            if name.startswith("target_"):
                cell.fill = target_fill
            elif name == "expected_area_mm2":
                cell.fill = predicted_fill

        hi = ws.cell(r, RUN_COLS.index("heat_input_j_mm") + 1)
        hi.value = (f'=IF(COUNT({col["actual_voltage_v"]}{r},{col["actual_current_a"]}{r},'
                    f'{col["actual_travel_mm_min"]}{r})=3,'
                    f'0.8*{col["actual_voltage_v"]}{r}*{col["actual_current_a"]}{r}*60/'
                    f'{col["actual_travel_mm_min"]}{r},"")')
        hi.fill = derived_fill


    ws = wb.create_sheet("Section Capture")
    hr = header(ws, SECTION_COLS,
                "White = measured (type it).  Green = formula (auto).  MID is mandatory; START/STOP = transient study.")
    sc = {name: get_column_letter(i) for i, name in enumerate(SECTION_COLS, 1)}
    r = hr
    for bead in beads():
        for loc in SECTIONS:
            r += 1
            seed = {"section_id": f'{bead["bead_id"]}-{loc}', "bead_id": bead["bead_id"], "location": loc,
                    "expected_area_mm2": expected_area_mm2(bead["wfs_ipm"], bead["travel_mm_min"])}
            for name in SECTION_COLS:
                cell = ws.cell(r, SECTION_COLS.index(name) + 1, seed.get(name))
                cell.border, cell.alignment = thin, wrap
                if name == "expected_area_mm2":
                    cell.fill = predicted_fill
            f = {"area_vs_expected": f'=IF(COUNT({sc["reinf_area_mm2"]}{r},{sc["expected_area_mm2"]}{r})=2,{sc["reinf_area_mm2"]}{r}/{sc["expected_area_mm2"]}{r},"")',
                 "total_fused_area_mm2": f'=IF(COUNT({sc["reinf_area_mm2"]}{r},{sc["penetration_area_mm2"]}{r})=2,{sc["reinf_area_mm2"]}{r}+{sc["penetration_area_mm2"]}{r},"")',
                 "dilution_pct": f'=IF(COUNT({sc["reinf_area_mm2"]}{r},{sc["penetration_area_mm2"]}{r})=2,{sc["penetration_area_mm2"]}{r}/({sc["reinf_area_mm2"]}{r}+{sc["penetration_area_mm2"]}{r})*100,"")',
                 "aspect_ratio_w_h": f'=IF(COUNT({sc["width_mm"]}{r},{sc["reinf_height_mm"]}{r})=2,{sc["width_mm"]}{r}/{sc["reinf_height_mm"]}{r},"")',
                 "parabola_area_mm2": f'=IF(COUNT({sc["width_mm"]}{r},{sc["reinf_height_mm"]}{r})=2,2/3*{sc["width_mm"]}{r}*{sc["reinf_height_mm"]}{r},"")',
                 "parabola_closure": f'=IF(COUNT({sc["reinf_area_mm2"]}{r},{sc["width_mm"]}{r},{sc["reinf_height_mm"]}{r})=3,{sc["reinf_area_mm2"]}{r}/(2/3*{sc["width_mm"]}{r}*{sc["reinf_height_mm"]}{r}),"")'}
            for name, formula in f.items():
                cell = ws.cell(r, SECTION_COLS.index(name) + 1, formula)
                cell.fill = derived_fill


    ws = wb.create_sheet("Definitions")
    header(ws, ["Quantity", "Definition"])
    for i, (q, d) in enumerate(DEFINITIONS, 2):
        ws.cell(i, 1, q).font = Font(bold=True)
        ws.cell(i, 1).alignment = wrap
        ws.cell(i, 2, d).alignment = wrap
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 110


    ws = wb.create_sheet("Macro Protocol")
    ws.cell(1, 1, "Sectioning + macro-etch protocol (ER70S-6 low-carbon steel)").font = Font(bold=True, size=12)
    for i, step in enumerate(PROTOCOL, 3):
        ws.cell(i, 1, step).alignment = wrap
    ws.column_dimensions["A"].width = 120


    for name, ws_name in (("Run Matrix", "Run Matrix"), ("Section Capture", "Section Capture")):
        sheet = wb[ws_name]
        for c in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(c)].width = 15
        sheet.column_dimensions[get_column_letter(1)].width = 16
    wb["Run Matrix"].column_dimensions[get_column_letter(len(RUN_COLS))].width = 34
    wb["Section Capture"].column_dimensions[get_column_letter(len(SECTION_COLS))].width = 34
    return wb


def section_csv() -> str:
    stream = io.StringIO(newline="")
    writer = csv.DictWriter(stream, fieldnames=SECTION_COLS)
    writer.writeheader()
    for bead in beads():
        for loc in SECTIONS:
            writer.writerow({"section_id": f'{bead["bead_id"]}-{loc}',
                             "bead_id": bead["bead_id"], "location": loc,
                             "expected_area_mm2": expected_area_mm2(bead["wfs_ipm"], bead["travel_mm_min"])})
    return stream.getvalue()


def run_matrix_csv() -> str:
    stream = io.StringIO(newline="")
    writer = csv.DictWriter(stream, fieldnames=RUN_COLS)
    writer.writeheader()
    for bead in beads():
        writer.writerow({"bead_id": bead["bead_id"], "plate": bead["plate"],
                         "target_wfs_ipm": bead["wfs_ipm"], "target_travel_mm_min": bead["travel_mm_min"],
                         "expected_area_mm2": expected_area_mm2(bead["wfs_ipm"], bead["travel_mm_min"]),
                         "arc_mode": PROC, "target_ctwd_mm": CTWD_MM, "gas": GAS,
                         "gas_flow_l_min": GAS_FLOW_L_MIN, "deposit_len_mm": DEPOSIT_LEN_MM,
                         "interpass_max_c": INTERPASS_MAX_C})
    return stream.getvalue()


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, required=True, help="output directory (created; must be new)")
    args = parser.parse_args(argv)
    try:
        if args.out.exists() and any(args.out.iterdir()):
            raise ValueError("Choose a new/empty output directory; existing records are preserved")
        gen_utc = datetime.now(timezone.utc).isoformat(timespec="seconds")
        wb = build_workbook(gen_utc)
        args.out.mkdir(parents=True, exist_ok=True)
        wb.save(args.out / "metallurgical-bead-capture.xlsx")
        (args.out / "metallurgical-bead-section-capture.csv").write_text(section_csv(), encoding="utf-8", newline="")
        (args.out / "metallurgical-bead-run-matrix.csv").write_text(run_matrix_csv(), encoding="utf-8", newline="")
        print(f"Capture pack: {args.out.resolve()}")
        print(f"  {len(beads())} beads x {len(SECTIONS)} sections = {len(beads()) * len(SECTIONS)} section rows")
        return 0
    except (OSError, ValueError, ImportError) as exc:
        print(f"Refused: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
