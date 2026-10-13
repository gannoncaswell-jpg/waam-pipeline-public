from __future__ import annotations

import csv
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
CSV = ROOT / "docs" / "measurement-sheet.csv"
XLSX = ROOT / "docs" / "measurement-sheet.xlsx"

INPUT_COLS = ("MEASURED_VALUE_MM", "MEASURED_BY", "DATE", "NOTES")


COLUMNS = [
    ("priority", "Pri", 5),
    ("component", "Component", 17),
    ("parameter", "What to measure", 40),
    ("how_to_measure", "How", 60),
    ("current_value_mm", "Model says (mm)", 20),
    ("basis", "Basis", 10),
    ("MEASURED_VALUE_MM", "MEASURED (mm)", 16),
    ("MEASURED_BY", "By", 10),
    ("DATE", "Date", 12),
    ("NOTES", "Notes", 30),
    ("id", "id (code key)", 24),
]

INK = "FF1B2430"
HDR_FILL = PatternFill("solid", fgColor="FF1F3350")
INPUT_FILL = PatternFill("solid", fgColor="FFFFF6D9")
DONE_FILL = PatternFill("solid", fgColor="FFE7F6EC")
PRI = {
    "H": PatternFill("solid", fgColor="FFFAD4D4"),
    "M": PatternFill("solid", fgColor="FFFDF0D0"),
    "L": PatternFill("solid", fgColor="FFEDEFF2"),
}
THIN = Side(style="thin", color="FFB9C2CC")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def existing_entries(path: Path) -> dict[str, dict[str, str]]:

    if not path.exists():
        return {}
    try:
        wb = load_workbook(path, data_only=True)
    except Exception as exc:
        raise SystemExit(f"cannot read existing {path.name} ({exc}); move it aside first")
    if "Measure" not in wb.sheetnames:
        return {}
    ws = wb["Measure"]
    head = [c.value for c in ws[1]]


    i_id = next((i for i, h in enumerate(head) if h in ("id (code key)", "id")), None)
    if i_id is None:
        raise SystemExit(
            f"{path.name} has no id column; refusing to overwrite it and lose data"
        )
    keep: dict[str, dict[str, str]] = {}
    labels = {lbl: key for key, lbl, _ in COLUMNS}
    labels["Measured value"] = "MEASURED_VALUE_MM"
    for row in ws.iter_rows(min_row=2):
        rid = row[i_id].value
        if not rid:
            continue
        vals = {}
        for idx, cell in enumerate(row):
            key = labels.get(head[idx]) if idx < len(head) else None
            if key in INPUT_COLS and cell.value not in (None, ""):
                vals[key] = str(cell.value).strip()
        if vals:
            keep[str(rid).strip()] = vals
    return keep


def main() -> int:
    if not CSV.exists():
        raise SystemExit(f"{CSV} missing - run gen_measurement_sheet.py first")
    rows = list(csv.DictReader(CSV.open(encoding="utf-8")))
    carried = existing_entries(XLSX)
    for r in rows:
        for k, v in carried.get(r["id"], {}).items():
            r[k] = v

    order = {"H": 0, "M": 1, "L": 2}
    rows.sort(key=lambda r: (order.get(r["priority"], 9), r["component"], r["id"]))

    wb = Workbook()
    ws = wb.active
    ws.title = "Measure"

    for i, (_key, label, width) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=i, value=label)
        c.font = Font(bold=True, color="FFFFFFFF", size=11)
        c.fill = HDR_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = BOX
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 30

    for r_i, row in enumerate(rows, start=2):
        for c_i, (key, _label, _w) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=r_i, column=c_i, value=row.get(key, ""))
            cell.border = BOX
            cell.alignment = Alignment(
                vertical="top", wrap_text=key in ("parameter", "how_to_measure", "NOTES")
            )
            if key in INPUT_COLS:
                cell.fill = INPUT_FILL
            elif key == "priority":
                cell.fill = PRI.get(row.get("priority", ""), PRI["L"])
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = Font(bold=True)
            elif key == "id":
                cell.font = Font(name="Consolas", size=9, color="FF6B7684")
            elif key == "current_value_mm":
                cell.font = Font(name="Consolas", size=10)
        ws.row_dimensions[r_i].height = 42

    last = len(rows) + 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last}"


    meas_col = get_column_letter(1 + [k for k, _, _ in COLUMNS].index("MEASURED_VALUE_MM"))
    ws.conditional_formatting.add(
        f"A2:{get_column_letter(len(COLUMNS))}{last}",
        FormulaRule(formula=[f'LEN($%s2)>0' % meas_col], fill=DONE_FILL, stopIfTrue=False),
    )

    dv = DataValidation(type="date", allow_blank=True,
                        error="Use a date, e.g. 2026-09-05", errorTitle="Date")
    ws.add_data_validation(dv)
    date_col = get_column_letter(1 + [k for k, _, _ in COLUMNS].index("DATE"))
    dv.add(f"{date_col}2:{date_col}{last}")


    g = wb.create_sheet("Read me")
    g.column_dimensions["A"].width = 108
    guide = [
        ("WAAM cell - measurements to gather", True),
        ("", False),
        ("Fill the four yellow columns. A row turns green once it has a value.", False),
        ("Sorted worst-first: do the red Pri=H rows before anything else.", False),
        ("", False),
        ("The 'id (code key)' column is the join back to the model - each id is a", False),
        ("constant in hardware/cell_models/cell_models.py or src/waam_pipeline/cell_lc.py.", False),
        ("Hand this back filled and the values go straight into the model.", False),
        ("", False),
        ("Basis = where the CURRENT number came from:", True),
        ("   MANUAL    a dimension stated in a vendor manual or print (best)", False),
        ("   MODEL     read out of a vendor CAD/simulator model", False),
        ("   LOG       a tape measurement or nameplate already recorded", False),
        ("   FK-RECON  back-solved from real robot positions in a pendant job", False),
        ("   SCALED    measured off a drawing figure (+/- a few mm)", False),
        ("   WEB       vendor spec page", False),
        ("   PHOTO     scaled off a lab photo - no scale bar, worst case", False),
        ("   ASSUMED   a guess that has never been checked  <- these are the risk", False),
        ("", False),
        ("Regenerating this file is safe: re-running the generator carries your", False),
        ("typed values across by id, so nothing you measured is lost.", False),
    ]
    for i, (text, bold) in enumerate(guide, start=1):
        c = g.cell(row=i, column=1, value=text)
        c.font = Font(bold=bold, size=13 if (bold and i == 1) else 11)
        c.alignment = Alignment(wrap_text=False, vertical="top")

    wb.save(XLSX)
    n_h = sum(1 for r in rows if r["priority"] == "H")
    print(f"wrote {XLSX.relative_to(ROOT)}  ({len(rows)} rows, {n_h} high priority, "
          f"{len(carried)} carried over)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
